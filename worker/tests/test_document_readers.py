"""Checks for real document formats entering the slow worker lane."""

import io
import unittest

from docx import Document
from PIL import Image, ImageDraw, ImageFont

from worker.services.document_parser import parse_documents


def _parse(content: bytes, media_type: str, filename: str) -> dict:
    result = parse_documents({
        "application_id": "application-1",
        "documents": [{"id": "document-1", "kind": "financial_statements",
                       "filename": filename, "media_type": media_type, "content": content}],
    })
    return result["documents"][0]


def _text_pdf() -> bytes:
    stream = b"BT /F1 18 Tf 50 750 Td (Credit Statement 2026) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Root 1 0 R /Size {len(offsets)} >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(pdf)


class DocumentReadersTest(unittest.TestCase):
    def test_docx_statement_table_extracts_reconciled_turnover(self):
        document = Document()
        table = document.add_table(rows=3, cols=4)
        for row, values in zip(table.rows, [
            ("Дата", "Поступление", "Списание", "Остаток"),
            ("02.01.2025", "100", "0", "1100"),
            ("03.01.2025", "0", "25", "1075"),
        ]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        output = io.BytesIO()
        document.save(output)
        item = parse_documents({"application_id": "application-1", "documents": [{
            "id": "statement-1", "kind": "account_statement", "filename": "statement.docx",
            "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "content": output.getvalue(),
        }]})["documents"][0]
        self.assertEqual(item["status"], "TEXT_EXTRACTED")
        fields = item["extracted_fields"]
        self.assertEqual(fields["turnover_monthly"]["value"][0]["credit"], "100.00")
        self.assertEqual(fields["closing_balance"]["value"], "1075.00")
        self.assertEqual(fields["turnover_monthly"]["source_document_id"], "statement-1")

    def test_docx_paragraphs_and_table(self):
        document = Document()
        document.add_paragraph("Revenue report 2026")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Revenue"
        table.cell(0, 1).text = "120000"
        output = io.BytesIO()
        document.save(output)

        item = _parse(output.getvalue(),
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                      "report.docx")
        self.assertEqual(item["status"], "TEXT_EXTRACTED")
        field = item["extracted_fields"]["document_text"]
        self.assertIn("Revenue\t120000", field["value"])
        self.assertEqual(field["source_document_id"], "document-1")
        self.assertEqual(field["verification_status"], "UNVERIFIED")

    def test_pdf_text_layer(self):
        item = _parse(_text_pdf(), "application/pdf", "report.pdf")
        self.assertEqual(item["status"], "TEXT_EXTRACTED")
        field = item["extracted_fields"]["document_text"]
        self.assertIn("Credit Statement 2026", field["value"])
        self.assertEqual(field["extraction_method"], "pdf_text")

    def test_image_ocr(self):
        image = Image.new("RGB", (1000, 180), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 44)
        draw.text((40, 50), "ACCOUNT 12345", fill="black", font=font)
        output = io.BytesIO()
        image.save(output, format="PNG")

        item = _parse(output.getvalue(), "image/png", "statement.png")
        self.assertEqual(item["status"], "TEXT_EXTRACTED")
        self.assertIn("ACCOUNT 12345", item["extracted_fields"]["document_text"]["value"])


if __name__ == "__main__":
    unittest.main()
